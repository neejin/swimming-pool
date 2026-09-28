"""
fetch_term_premia.py — NY Fed ACM 10Y (daily) + FRED DFII10 → data/us10y_decomposition.xlsx

Step 1. NY Fed ACMTermPremium.xls 의 daily 시트에서 ACMTP10, ACMRNY10, ACMY10 추출
Step 2. FRED DFII10(daily)을 날짜 기준으로 merge (outer join, 한쪽만 있는 날짜는 빈칸)

매 실행마다 두 소스 전체 이력을 새로 받아 파일을 통째로 다시 만든다
(ACM 추정치가 과거 구간까지 수정될 수 있으므로 append 방식 대신 전체 재생성).
"""
import io, os
from datetime import datetime, timezone

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

ACM_URL      = "https://www.newyorkfed.org/medialibrary/media/research/data_indicators/ACMTermPremium.xls"
ACM_PAGE     = "https://www.newyorkfed.org/research/data_indicators/term-premia-tabs"
FRED_PAGE    = "https://fred.stlouisfed.org/series/DFII10"
FRED_CSV     = "https://fred.stlouisfed.org/graph/fredgraph.csv?id=DFII10"
FRED_API_KEY = os.environ.get("FRED_API_KEY", "")
ACM_COLS     = ["ACMTP10", "ACMRNY10", "ACMY10"]
OUT          = "data/us10y_decomposition.xlsx"
HEADERS      = {"User-Agent": "Mozilla/5.0 (data fetch script)"}


def read_acm_daily(content):
    # .xls(OLE2)인지 .xlsx(zip)인지 매직 바이트로 판별
    engine = "openpyxl" if content[:2] == b"PK" else "xlrd"
    xls = pd.ExcelFile(io.BytesIO(content), engine=engine)
    print(f"[ACM] sheets: {xls.sheet_names}")
    sheet = next((s for s in xls.sheet_names if "daily" in s.lower()), None)
    if sheet is None:
        raise RuntimeError(f"daily 시트를 찾지 못함: {xls.sheet_names}")

    raw = xls.parse(sheet, header=None)
    # 헤더 행 = ACMTP10 이 들어있는 첫 행
    hdr = next(i for i, row in raw.iterrows()
               if any(str(v).strip().upper() == "ACMTP10" for v in row.values))
    df = raw.iloc[hdr + 1:].copy()
    df.columns = [str(c).strip().upper() for c in raw.iloc[hdr]]
    missing = [c for c in ACM_COLS if c not in df.columns]
    if missing:
        raise RuntimeError(f"[ACM] 컬럼 없음: {missing}")

    date_col = "DATE" if "DATE" in df.columns else df.columns[0]
    out = pd.DataFrame({"Date": pd.to_datetime(df[date_col], errors="coerce")})
    for c in ACM_COLS:
        out[c] = pd.to_numeric(df[c], errors="coerce")
    out = out.dropna(subset=["Date"]).dropna(subset=ACM_COLS, how="all")
    out = out.drop_duplicates("Date", keep="last").sort_values("Date")
    print(f"[ACM] {sheet}: {len(out):,} rows, "
          f"{out['Date'].min():%Y-%m-%d} ~ {out['Date'].max():%Y-%m-%d}")
    return out, sheet


def fetch_acm():
    print(f"[ACM] GET {ACM_URL}")
    r = requests.get(ACM_URL, headers=HEADERS, timeout=120)
    r.raise_for_status()
    print(f"[ACM] HTTP {r.status_code}, {len(r.content):,} bytes")
    return read_acm_daily(r.content)


def fetch_dfii10():
    if FRED_API_KEY:
        url = ("https://api.stlouisfed.org/fred/series/observations"
               f"?series_id=DFII10&api_key={FRED_API_KEY}&file_type=json&sort_order=asc")
        print("[FRED] GET api.stlouisfed.org DFII10 (API)")
        r = requests.get(url, timeout=60)
        r.raise_for_status()
        obs = r.json()["observations"]
        df = pd.DataFrame({"Date": [o["date"] for o in obs],
                           "DFII10": [o["value"] for o in obs]})
    else:
        print(f"[FRED] GET {FRED_CSV} (no API key)")
        r = requests.get(FRED_CSV, headers=HEADERS, timeout=60)
        r.raise_for_status()
        df = pd.read_csv(io.StringIO(r.text))
        df.columns = ["Date", "DFII10"]  # observation_date(또는 DATE), DFII10
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df["DFII10"] = pd.to_numeric(df["DFII10"], errors="coerce")  # "." / 빈칸 → NaN
    df = df.dropna().sort_values("Date")
    print(f"[FRED] DFII10: {len(df):,} rows, "
          f"{df['Date'].min():%Y-%m-%d} ~ {df['Date'].max():%Y-%m-%d}")
    return df


def write_xlsx(df, acm_sheet):
    font  = Font(name="Arial", size=10)
    bold  = Font(name="Arial", size=10, bold=True)
    fill  = PatternFill("solid", fgColor="D9E1F2")
    wb = Workbook()

    ws = wb.active
    ws.title = "Data"
    cols = ["Date"] + ACM_COLS + ["DFII10"]
    for j, name in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=name)
        c.font, c.fill, c.alignment = bold, fill, Alignment(horizontal="center")
    for i, row in enumerate(df[cols].itertuples(index=False), 2):
        ws.cell(row=i, column=1, value=row[0].to_pydatetime().date()).number_format = "yyyy-mm-dd"
        for j, v in enumerate(row[1:], 2):
            if pd.notna(v):
                ws.cell(row=i, column=j, value=float(v)).number_format = "0.0000"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = font
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 12
    for j in range(2, len(cols) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 11

    info = wb.create_sheet("Info")
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    last = lambda c: df.loc[df[c].notna(), "Date"].max()
    rows = [
        ("Updated", now),
        ("Unit", "percent (%)"),
        ("Merge", "outer join on Date (blank = not available from that source)"),
        ("", ""),
        ("Column", "Description / Source"),
        ("ACMTP10", f"ACM 10Y term premium — NY Fed, sheet '{acm_sheet}' — {ACM_PAGE}"),
        ("ACMRNY10", f"ACM 10Y risk-neutral yield — NY Fed — {ACM_PAGE}"),
        ("ACMY10", f"ACM 10Y fitted zero-coupon yield — NY Fed — {ACM_PAGE}"),
        ("DFII10", f"10Y TIPS constant-maturity real yield (daily) — FRED — {FRED_PAGE}"),
        ("", ""),
        ("Last date ACM", f"{last('ACMY10'):%Y-%m-%d}"),
        ("Last date DFII10", f"{last('DFII10'):%Y-%m-%d}"),
    ]
    for i, (k, v) in enumerate(rows, 1):
        info.cell(row=i, column=1, value=k).font = bold
        info.cell(row=i, column=2, value=v).font = font
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 100

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT} — {len(df):,} rows")


def main():
    acm, acm_sheet = fetch_acm()
    dfii = fetch_dfii10()
    merged = acm.merge(dfii, on="Date", how="outer").sort_values("Date")
    write_xlsx(merged.reset_index(drop=True), acm_sheet)


if __name__ == "__main__":
    main()
